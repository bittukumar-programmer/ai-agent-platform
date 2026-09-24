import os
import subprocess
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

# Base folder where the assistant will create files/folders, to keep things safe and organized
WORKSPACE = os.path.join(os.path.expanduser("~"), "Desktop", "AI_Workspace")
os.makedirs(WORKSPACE, exist_ok=True)


def open_notepad(filename: str = None) -> str:
    """Opens Notepad, optionally with a specific file."""
    try:
        if filename:
            path = os.path.join(WORKSPACE, filename)
            if not os.path.exists(path):
                open(path, "w").close()
            subprocess.Popen(["notepad.exe", path])
            return f"Opened Notepad with {filename}"
        else:
            subprocess.Popen(["notepad.exe"])
            return "Opened Notepad"
    except Exception as e:
        return f"Error opening Notepad: {e}"


def create_folder(folder_name: str) -> str:
    """Creates a new folder inside the AI Workspace."""
    try:
        path = os.path.join(WORKSPACE, folder_name)
        os.makedirs(path, exist_ok=True)
        return f"Created folder '{folder_name}' at {path}"
    except Exception as e:
        return f"Error creating folder: {e}"


def create_file(filename: str, folder_name: str = None, content: str = "") -> str:
    """Creates a new file, optionally inside a specific folder within the workspace."""
    try:
        base = os.path.join(WORKSPACE, folder_name) if folder_name else WORKSPACE
        os.makedirs(base, exist_ok=True)
        path = os.path.join(base, filename)
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return f"Created file '{filename}' at {path}"
    except Exception as e:
        return f"Error creating file: {e}"


def open_vscode(path: str = None) -> str:
    """Opens VS Code, optionally at a specific folder/file path (relative to the workspace)."""
    try:
        if path:
            # If it's already a full path, use it directly; otherwise treat it as relative to workspace
            target = path if os.path.isabs(path) else os.path.join(WORKSPACE, path)
        else:
            target = WORKSPACE
        subprocess.Popen(["code", target], shell=True)
        return f"Opened VS Code at {target}"
    except Exception as e:
        return f"Error opening VS Code: {e}"


def close_app(app_name: str) -> str:
    """Force-closes an application by its process name (e.g. 'Code.exe', 'notepad.exe')."""
    try:
        subprocess.run(["taskkill", "/IM", app_name, "/F"], capture_output=True)
        return f"Closed {app_name}"
    except Exception as e:
        return f"Error closing {app_name}: {e}"


def open_app(app_name: str) -> str:
    """Opens any application by its common name (Chrome, Calculator, WhatsApp, etc.)."""
    # Common app name → actual Windows 'start' command mapping
    app_map = {
        "chrome": "chrome.exe",
        "google chrome": "chrome.exe",
        "calculator": "calc.exe",
        "calc": "calc.exe",
        "paint": "mspaint.exe",
        "explorer": "explorer.exe",
        "file explorer": "explorer.exe",
        "whatsapp": "shell:AppsFolder\\5319275A.WhatsAppDesktop_cv1g1gvanyjgm!App",
        "word": "winword.exe",
        "excel": "excel.exe",
        "powerpoint": "powerpnt.exe",
        "spotify": "spotify.exe",
        "notepad": "notepad.exe",
    }
    try:
        key = app_name.lower().strip()
        command = app_map.get(key, app_name)
        # 'start' is Windows's own reliable way to launch apps registered on the system
        subprocess.Popen(f'start "" "{command}"', shell=True)
        return f"Opened {app_name}"
    except Exception as e:
        return f"Error opening {app_name}: {e}"


def open_folder_in_explorer(folder_path: str) -> str:
    """Opens a folder in Windows File Explorer (relative to the AI workspace, or absolute)."""
    try:
        target = folder_path if os.path.isabs(folder_path) else os.path.join(WORKSPACE, folder_path)
        os.makedirs(target, exist_ok=True)
        subprocess.Popen(["explorer", target])
        return f"Opened {target} in File Explorer"
    except Exception as e:
        return f"Error opening folder: {e}"


def search_web(query: str) -> str:
    """Opens the default browser with a Google search for the given query."""
    try:
        import webbrowser
        url = f"https://www.google.com/search?q={query.replace(' ', '+')}"
        webbrowser.open(url)
        return f"Searched for '{query}' in your browser"
    except Exception as e:
        return f"Error searching: {e}"



def create_word_doc(filename: str, title: str, content: str, folder_name: str = None) -> str:
    """Creates a Word document with a title and body content."""
    try:
        base = os.path.join(WORKSPACE, folder_name) if folder_name else WORKSPACE
        os.makedirs(base, exist_ok=True)
        if not filename.endswith(".docx"):
            filename += ".docx"
        path = os.path.join(base, filename)

        doc = Document()
        doc.add_heading(title, level=1)
        for paragraph in content.split("\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph)
        doc.save(path)
        return f"Created Word document '{filename}' at {path}"
    except Exception as e:
        return f"Error creating Word document: {e}"


def create_excel_sheet(filename: str, sheet_title: str, headers: list, rows: list, folder_name: str = None) -> str:
    """Creates an Excel sheet with headers and data rows."""
    try:
        base = os.path.join(WORKSPACE, folder_name) if folder_name else WORKSPACE
        os.makedirs(base, exist_ok=True)
        if not filename.endswith(".xlsx"):
            filename += ".xlsx"
        path = os.path.join(base, filename)

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title[:31]  # Excel sheet name limit
        if headers:
            ws.append(headers)
        for row in rows:
            ws.append(row)
        wb.save(path)
        return f"Created Excel sheet '{filename}' at {path}"
    except Exception as e:
        return f"Error creating Excel sheet: {e}"


def create_powerpoint(filename: str, slides: list, folder_name: str = None) -> str:
    """Creates a PowerPoint presentation. slides = list of {'title': ..., 'content': ...}"""
    try:
        base = os.path.join(WORKSPACE, folder_name) if folder_name else WORKSPACE
        os.makedirs(base, exist_ok=True)
        if not filename.endswith(".pptx"):
            filename += ".pptx"
        path = os.path.join(base, filename)

        prs = Presentation()
        title_layout = prs.slide_layouts[1]  # title + content layout

        for slide_data in slides:
            slide = prs.slides.add_slide(title_layout)
            slide.shapes.title.text = slide_data.get("title", "")
            body = slide.placeholders[1]
            body.text = slide_data.get("content", "")

        prs.save(path)
        return f"Created PowerPoint '{filename}' with {len(slides)} slides at {path}"
    except Exception as e:
        return f"Error creating PowerPoint: {e}"

def open_file(file_path: str) -> str:
    """Opens any file with its default application (e.g. a .pptx opens in PowerPoint)."""
    try:
        target = file_path if os.path.isabs(file_path) else os.path.join(WORKSPACE, file_path)
        os.startfile(target)
        return f"Opened {target}"
    except Exception as e:
        return f"Error opening file: {e}"